from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as L
from datetime import datetime as D, timedelta as T

wb = Workbook()
FN = 'Arial'
hdr_fill = PatternFill('solid', fgColor='17202A'); hdr_font = Font(name=FN, bold=True, color='FFFFFF', size=10)
inp = Font(name=FN, color='0000FF', size=10); calc = Font(name=FN, color='000000', size=10)
bold = Font(name=FN, bold=True, size=10); title = Font(name=FN, bold=True, size=14, color='0F766E')
note = Font(name=FN, italic=True, size=9, color='595959')
yellow = PatternFill('solid', fgColor='FFF2CC'); grey = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='D9D9D9'); box = Border(left=thin, right=thin, top=thin, bottom=thin)
green = PatternFill('solid', bgColor='C6EFCE', fgColor='C6EFCE'); red = PatternFill('solid', bgColor='FFC7CE', fgColor='FFC7CE')

def head(ws, row, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row, i, c); x.font = hdr_font; x.fill = hdr_fill; x.border = box
        x.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[row].height = 32
def put(ws, r, c, v, kind='calc', fmt=None):
    x = ws.cell(r, c, v); x.font = inp if kind == 'inp' else calc; x.border = box
    if fmt: x.number_format = fmt
    return x
def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1): ws.column_dimensions[L(i)].width = w
def flag(ws, rng, good, bad):
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=[f'"{good}"'], fill=green))
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=[f'"{bad}"'], fill=red))
DT = 'yyyy-mm-dd hh:mm'

# ---------- README
ws = wb.active; ws.title = 'README'
ws['A1'] = 'HR Talent Acquisition & Onboarding Workflow: Working Model'; ws['A1'].font = title
rows = [
 ('Purpose', 'A working spreadsheet version of the Week 1 to 4 framework: ATS tracker, scorecards, onboarding tracker, SLA and pilot results, and a KPI dashboard. All formulas recalculate when you edit inputs.'),
 ('Data', 'All names, emails and figures are dummy data from the simulated case (Northwind Services Pvt. Ltd.). Hire sample = Week 1 audit sample; pilot counts = Week 3 pilot.'),
 ('Legend', 'Blue text = input you can edit. Black text = formula, do not overwrite. Yellow fill = setting that drives formulas (change it on the Settings sheet). Green/red = status flags.'),
 ('Candidates', 'ATS tracker. Add a row per applicant (rows 5 to 44 are formula-ready). Stage and Source use dropdowns. Duplicate check flags repeated emails; acknowledgement SLA and screening result calculate automatically.'),
 ('Scorecard', 'One row per interviewer scorecard. Enter ratings 1 to 5. Weighted score, feedback hours, SLA status, decision guide and debrief gap flag calculate automatically.'),
 ('Hires_Sample', '10 past hires from the Week 1 audit. Gives the 45-day time-to-hire and 70% retention baselines.'),
 ('Onboarding', 'Tracker for new hires: checklist, IT access, LMS modules, Day 10 reminder flag and completion %.'),
 ('Pilot_Results', 'The 12 pilot metrics from Week 3. Pilot score = metrics at target divided by total (target 80%).'),
 ('KPI_Dashboard', 'Live KPI view with status flags, pipeline funnel and charts. Pulls from all other sheets.'),
 ('Settings', 'Thresholds (screening cut-off, 24-hour SLA, debrief gap) and the scorecard criteria weights.'),
]
for i, (a, b) in enumerate(rows, 3):
    ws.cell(i, 1, a).font = bold; c = ws.cell(i, 2, b); c.font = calc; c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[i].height = 44
widths(ws, [18, 110])

# ---------- Settings
st = wb.create_sheet('Settings')
st['A1'] = 'Settings'; st['A1'].font = title
setr = [('Screening pass score (out of 100)', 60), ('Scorecard SLA (hours after interview)', 24),
        ('Acknowledgement SLA (hours)', 24), ('Debrief gap flag (rating points)', 1.5),
        ('Strong hire threshold (weighted score)', 4.0), ('Discuss threshold (weighted score)', 3.0),
        ('Screening SLA (working days)', 3)]
for i, (a, b) in enumerate(setr, 2):
    st.cell(i, 1, a).font = calc; x = st.cell(i, 2, b); x.font = inp; x.fill = yellow; x.border = box
st['A10'] = 'Scorecard criteria'; st['A10'].font = bold
crit = ['Role skills', 'Problem solving', 'Communication', 'Values alignment', 'Learning ability']; w = [.30, .20, .20, .15, .15]
st['A11'] = 'Criterion'; st['A12'] = 'Weight'
for j, (c, wt) in enumerate(zip(crit, w), 2):
    st.cell(11, j, c).font = bold; x = st.cell(12, j, wt); x.font = inp; x.fill = yellow; x.number_format = '0%'; x.border = box
st['A11'].font = bold; st['A12'].font = calc
st['G11'] = 'Total'; st['G11'].font = bold; st['G12'] = '=SUM(B12:F12)'; st['G12'].number_format = '0%'; st['G12'].font = calc
st['H12'] = '=IF(ROUND(G12,4)=1,"OK","Weights must total 100%")'; st['H12'].font = calc
st['A14'] = 'Source: thresholds and weights from Week 2 templates T-03 and T-04 and Week 3 SLA register.'; st['A14'].font = note
widths(st, [42, 14, 16, 14, 16, 14, 10, 24])
SET = dict(screen='Settings!$B$2', fb='Settings!$B$3', ack='Settings!$B$4', gap='Settings!$B$5', strong='Settings!$B$6', disc='Settings!$B$7', wd='Settings!$B$8', w='Settings!$B$12:$F$12')

# ---------- Candidates
ca = wb.create_sheet('Candidates')
ca['A1'] = 'Candidate Tracker (ATS)'; ca['A1'].font = title
ca['A2'] = 'Blue = input. Add rows below the sample data (formulas are ready to row 44). Duplicate emails are flagged.'; ca['A2'].font = note
cols = ['Cand ID', 'Name', 'Email', 'Role', 'Req ID', 'Source', 'Stage', 'Applied', 'Acknowledged', 'Ack hours', 'Ack SLA', 'Screen score', 'Screening result', 'Duplicate check', 'Rejection reason']
head(ca, 4, cols)
base = D(2026, 10, 12, 9, 0)
data = [
 ('Candidate A', 'candidate.a@example.com', 'Customer Support Agent', 'REQ-101', 'LinkedIn', 'Interview', 0, 3, 78, ''),
 ('Candidate B', 'candidate.b@example.com', 'Customer Support Agent', 'REQ-101', 'Job board', 'Screening', 5, 5, 66, ''),
 ('Candidate C', 'candidate.c@example.com', 'Customer Support Agent', 'REQ-101', 'Referral', 'Rejected', 10, 2, 48, 'Below cut-off score'),
 ('Candidate D', 'candidate.d@example.com', 'Customer Support Agent', 'REQ-101', 'Email', 'Screening', 20, 22, 71, ''),
 ('Candidate E', 'candidate.e@example.com', 'Customer Support Agent', 'REQ-101', 'Job board', 'Applied', 26, 4, None, ''),
 ('Candidate F', 'candidate.f@example.com', 'Customer Support Agent', 'REQ-101', 'LinkedIn', 'Applied', 30, 30, None, ''),
 ('Candidate G', 'candidate.g@example.com', 'Marketing Executive', 'REQ-102', 'LinkedIn', 'Interview', 8, 1, 82, ''),
 ('Candidate H', 'candidate.h@example.com', 'Sales Associate', 'REQ-103', 'Referral', 'Offer', 12, 6, 90, ''),
 ('Candidate I', 'candidate.i@example.com', 'Marketing Executive', 'REQ-102', 'Job board', 'Hired', 15, 12, 85, ''),
 ('Candidate J', 'candidate.j@example.com', 'Sales Associate', 'REQ-103', 'Email', 'Decision', 18, 2, 74, ''),
 ('Candidate A', 'candidate.a@example.com', 'Customer Support Agent', 'REQ-101', 'Job board', 'Applied', 40, 5, None, ''),
 ('Candidate D', 'candidate.d@example.com', 'Customer Support Agent', 'REQ-101', 'Email', 'Applied', 44, 3, None, ''),
]
LAST = 44
for r in range(5, LAST + 1):
    i = r - 5
    if i < len(data):
        n, em, role, req, src, stg, off, ackh, sc, rej = data[i]
        ap = base + T(hours=off * 3)
        put(ca, r, 1, f'CAN-{i+1:03d}', 'inp'); put(ca, r, 2, n, 'inp'); put(ca, r, 3, em, 'inp'); put(ca, r, 4, role, 'inp')
        put(ca, r, 5, req, 'inp'); put(ca, r, 6, src, 'inp'); put(ca, r, 7, stg, 'inp')
        put(ca, r, 8, ap, 'inp', DT); put(ca, r, 9, ap + T(hours=ackh), 'inp', DT); put(ca, r, 12, sc, 'inp'); put(ca, r, 15, rej, 'inp')
    else:
        for c in (1,2,3,4,5,6,7,8,9,12,15): put(ca, r, c, None, 'inp', DT if c in (8,9) else None)
    put(ca, r, 10, f'=IF(OR(H{r}="",I{r}=""),"",ROUND((I{r}-H{r})*24,1))', fmt='0.0')
    put(ca, r, 11, f'=IF(J{r}="","",IF(J{r}<={SET["ack"]},"On time","Late"))')
    put(ca, r, 13, f'=IF(L{r}="","",IF(L{r}>={SET["screen"]},"Shortlist","Below cut-off"))')
    put(ca, r, 14, f'=IF(C{r}="","",IF(COUNTIF($C$5:$C${LAST},C{r})>1,"Duplicate","OK"))')
stages = ['Applied', 'Screening', 'Interview', 'Decision', 'Offer', 'Hired', 'Rejected']
dv = DataValidation(type='list', formula1='"' + ','.join(stages) + '"', allow_blank=True); ca.add_data_validation(dv); dv.add(f'G5:G{LAST}')
dv2 = DataValidation(type='list', formula1='"LinkedIn,Job board,Referral,Email,Walk-in"', allow_blank=False); ca.add_data_validation(dv2); dv2.add(f'F5:F{LAST}')
flag(ca, f'K5:K{LAST}', 'On time', 'Late'); flag(ca, f'N5:N{LAST}', 'OK', 'Duplicate'); flag(ca, f'M5:M{LAST}', 'Shortlist', 'Below cut-off')
ca.freeze_panes = 'C5'; widths(ca, [10, 15, 28, 22, 9, 11, 11, 17, 17, 9, 10, 10, 15, 14, 22])

# ---------- Scorecard
sc = wb.create_sheet('Scorecard')
sc['A1'] = 'Interview Scorecards'; sc['A1'].font = title
sc['A2'] = 'Ratings 1 to 5 (1 = does not meet, 3 = meets, 5 = outstanding). Weights and thresholds come from Settings.'; sc['A2'].font = note
head(sc, 4, ['Interview ID', 'Cand ID', 'Candidate', 'Interviewer', 'Interview end', 'Scorecard submitted'] + crit + ['Evidence on every criterion', 'Weighted score', 'Feedback hours', 'Feedback SLA', 'Decision guide', 'Rating gap (same candidate)', 'Debrief flag'])
SL = 40
iv = ['Sara Malik', 'Usman Tariq', 'Bilal Ahmed', 'Ayesha Khan']
cand = ['CAN-007', 'CAN-007', 'CAN-009', 'CAN-009', 'CAN-001', 'CAN-001', 'CAN-010', 'CAN-010', 'CAN-008', 'CAN-008']
hrs = [4, 18, 22, 6, 29, 20, 12, 31, 15, 23]
rat = [(4,4,5,4,4),(4,4,4,5,4),(5,4,4,4,5),(4,5,4,4,4),(3,3,4,4,3),(4,3,3,4,4),(4,4,4,3,4),(3,3,3,4,3),(5,5,4,5,4),(4,5,5,4,5)]
for r in range(5, SL + 1):
    i = r - 5
    if i < 10:
        end = D(2026, 10, 20, 10, 0) + T(hours=i * 5); sub = end + T(hours=hrs[i])
        put(sc, r, 1, f'INT-{i+1:03d}', 'inp'); put(sc, r, 2, cand[i], 'inp'); put(sc, r, 4, iv[i % 4], 'inp')
        put(sc, r, 5, end, 'inp', DT); put(sc, r, 6, sub, 'inp', DT)
        for j in range(5): put(sc, r, 7 + j, rat[i][j], 'inp')
        put(sc, r, 12, 'No' if i == 4 else 'Yes', 'inp')
    else:
        for c in list(range(1, 3)) + list(range(4, 13)): put(sc, r, c, None, 'inp', DT if c in (5, 6) else None)
    put(sc, r, 3, f'=IF(B{r}="","",IFERROR(INDEX(Candidates!$B$5:$B${LAST},MATCH(B{r},Candidates!$A$5:$A${LAST},0)),"Not found"))')
    put(sc, r, 13, f'=IF(COUNT(G{r}:K{r})=0,"",SUMPRODUCT(G{r}:K{r},{SET["w"]})/SUMPRODUCT((G{r}:K{r}<>"")*{SET["w"]}))', fmt='0.00')
    put(sc, r, 14, f'=IF(OR(E{r}="",F{r}=""),"",ROUND((F{r}-E{r})*24,1))', fmt='0.0')
    put(sc, r, 15, f'=IF(N{r}="","",IF(N{r}<={SET["fb"]},"On time","Late"))')
    put(sc, r, 16, f'=IF(M{r}="","",IF(M{r}>={SET["strong"]},"Strong hire",IF(M{r}>={SET["disc"]},"Discuss at debrief","Do not proceed")))')
    put(sc, r, 17, f'=IF(M{r}="","",_xlfn.MAXIFS($M$5:$M${SL},$B$5:$B${SL},B{r})-_xlfn.MINIFS($M$5:$M${SL},$B$5:$B${SL},B{r}))', fmt='0.00')
    put(sc, r, 18, f'=IF(Q{r}="","",IF(Q{r}>{SET["gap"]},"Debrief gap","OK"))')
dv3 = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True, showErrorMessage=True, error='Enter a rating from 1 to 5'); sc.add_data_validation(dv3); dv3.add(f'G5:K{SL}')
dv4 = DataValidation(type='list', formula1='"Yes,No"', allow_blank=True); sc.add_data_validation(dv4); dv4.add(f'L5:L{SL}')
flag(sc, f'O5:O{SL}', 'On time', 'Late'); flag(sc, f'R5:R{SL}', 'OK', 'Debrief gap')
sc.freeze_panes = 'D5'; widths(sc, [12, 10, 15, 14, 17, 17, 9, 11, 13, 11, 11, 13, 10, 10, 11, 19, 12, 12])

# ---------- Hires_Sample
hs = wb.create_sheet('Hires_Sample')
hs['A1'] = 'Hires Sample (Week 1 baseline)'; hs['A1'].font = title
hs['A2'] = 'Source: 10 past hires from the simulated Week 1 audit (requisitions Jan to May 2025). Illustrative data.'; hs['A2'].font = note
head(hs, 4, ['Req ID', 'Role', 'Request date', 'Offer accepted', 'Days-to-hire', 'Source', 'Retained at day 90'])
hires = [('H-01','Marketing Executive','2025-01-12','2025-03-05','LinkedIn','Yes'),('H-02','Software Developer','2025-01-26','2025-03-05','Referral','Yes'),
 ('H-03','Sales Associate','2025-02-09','2025-03-28','Job board','No'),('H-04','Accountant','2025-02-23','2025-04-05','LinkedIn','Yes'),
 ('H-05','Operations Analyst','2025-03-02','2025-04-26','Job board','Yes'),('H-06','Customer Support Agent','2025-03-16','2025-04-21','Email','No'),
 ('H-07','HR Assistant','2025-04-06','2025-05-25','Referral','Yes'),('H-08','Graphic Designer','2025-04-20','2025-06-02','LinkedIn','Yes'),
 ('H-09','Sales Associate','2025-05-04','2025-06-06','Job board','No'),('H-10','Customer Support Agent','2025-05-18','2025-07-13','Email','Yes')]
for i, h in enumerate(hires):
    r = 5 + i
    put(hs, r, 1, h[0], 'inp'); put(hs, r, 2, h[1], 'inp'); put(hs, r, 3, D.strptime(h[2], '%Y-%m-%d'), 'inp', 'yyyy-mm-dd'); put(hs, r, 4, D.strptime(h[3], '%Y-%m-%d'), 'inp', 'yyyy-mm-dd')
    put(hs, r, 5, f'=D{r}-C{r}', fmt='0'); put(hs, r, 6, h[4], 'inp'); put(hs, r, 7, h[5], 'inp')
hs['A16'] = 'Average days-to-hire'; hs['E16'] = '=AVERAGE(E5:E14)'; hs['E16'].number_format = '0.0'
hs['A17'] = '90-day retention'; hs['E17'] = '=COUNTIF(G5:G14,"Yes")/COUNTA(G5:G14)'; hs['E17'].number_format = '0%'
hs['A18'] = 'Hires in sample'; hs['E18'] = '=COUNTA(A5:A14)'
for r in (16, 17, 18): hs.cell(r, 1).font = bold; hs.cell(r, 5).font = bold
flag(hs, 'G5:G14', 'Yes', 'No'); widths(hs, [10, 24, 14, 15, 13, 12, 17])

# ---------- Onboarding
ob = wb.create_sheet('Onboarding')
ob['A1'] = 'Onboarding Tracker'; ob['A1'].font = title
ob['A2'] = 'Day 10 snapshot for the 3 pilot hires (dummy). 5 mandatory LMS modules each. Reminder rule R8 fires on Day 10 if modules are incomplete.'; ob['A2'].font = note
head(ob, 4, ['New hire', 'Role', 'Start date', 'Current day', 'Checklist created', 'IT access ready Day 1', 'LMS enrolled', 'Modules done', 'Modules assigned', 'LMS completion', 'Reminder (Day 10)', 'Day 14 status'])
nh = [('New hire A', 'Customer Support Agent', 5), ('New hire B', 'Marketing Executive', 4), ('New hire C', 'Sales Associate', 5)]
for i, (n, ro, m) in enumerate(nh):
    r = 5 + i
    put(ob, r, 1, n, 'inp'); put(ob, r, 2, ro, 'inp'); put(ob, r, 3, D(2026, 10, 19), 'inp', 'yyyy-mm-dd'); put(ob, r, 4, 10, 'inp')
    for c in (5, 6, 7): put(ob, r, c, 'Yes', 'inp')
    put(ob, r, 8, m, 'inp'); put(ob, r, 9, 5, 'inp')
    put(ob, r, 10, f'=IF(I{r}=0,"",H{r}/I{r})', fmt='0%')
    put(ob, r, 11, f'=IF(AND(D{r}>=10,H{r}<I{r}),"Send reminder","OK")')
    put(ob, r, 12, f'=IF(H{r}>=I{r},"Complete",IF(D{r}>=14,"Overdue","In progress"))')
for r in range(8, 15):
    for c in range(1, 10): put(ob, r, c, None, 'inp')
    put(ob, r, 10, f'=IF(OR(I{r}="",I{r}=0),"",H{r}/I{r})', fmt='0%'); put(ob, r, 11, f'=IF(A{r}="","",IF(AND(D{r}>=10,H{r}<I{r}),"Send reminder","OK"))'); put(ob, r, 12, f'=IF(A{r}="","",IF(H{r}>=I{r},"Complete",IF(D{r}>=14,"Overdue","In progress")))')
dv5 = DataValidation(type='list', formula1='"Yes,No"', allow_blank=True); ob.add_data_validation(dv5); dv5.add('E5:G14')
flag(ob, 'K5:K14', 'OK', 'Send reminder'); widths(ob, [14, 24, 12, 11, 12, 14, 11, 11, 12, 12, 16, 14])
ob['A16'] = 'Average LMS completion'; ob['A16'].font = bold; ob['J16'] = '=AVERAGE(J5:J14)'; ob['J16'].number_format = '0.0%'; ob['J16'].font = bold

# ---------- Pilot_Results
pr = wb.create_sheet('Pilot_Results')
pr['A1'] = 'Pilot Results (Week 3)'; pr['A1'].font = title
pr['A2'] = 'Counts are from the simulated 5-day pilot. P3 first read was 17 of 21 (81.0%) before the working-day formula fix; the corrected 19 of 21 is shown.'; pr['A2'].font = note
head(pr, 4, ['ID', 'Metric', 'Target', 'Sample (n)', 'On target (count)', 'Result', 'Met target?', 'Counts toward SLA total?'])
pm = [('P1','Application acknowledged within 24 hours',.95,24,23,'Yes'),('P2','Seeded duplicates flagged',1,3,3,'No'),('P3','Screening within 3 working days',.90,21,19,'Yes'),
 ('P4','Scorecards submitted within 24 hours',.90,10,8,'Yes'),('P5','Scorecards complete (evidence on every criterion)',.90,10,9,'No'),('P6','Decision within 2 working days',.90,2,2,'Yes'),
 ('P7','Offer sent within 1 working day',.90,2,2,'Yes'),('P8','Offer signed within 3 working days',.80,2,1,'Yes'),('P9','IT access ready on Day 1',1,3,3,'Yes'),
 ('P10','Onboarding checklist created automatically',1,3,3,'No'),('P11','LMS modules completed by Day 5 (interim, module slots)',.60,15,11,'No'),('P12','Candidate satisfaction (4 or 5 of 5)',.85,20,17,'No')]
for i, p in enumerate(pm):
    r = 5 + i
    put(pr, r, 1, p[0], 'inp'); put(pr, r, 2, p[1], 'inp'); put(pr, r, 3, p[2], 'inp', '0%'); put(pr, r, 4, p[3], 'inp'); put(pr, r, 5, p[4], 'inp')
    put(pr, r, 6, f'=IF(D{r}=0,"",E{r}/D{r})', fmt='0.0%'); put(pr, r, 7, f'=IF(F{r}="","",IF(F{r}>=C{r},"Yes","No"))'); put(pr, r, 8, p[5], 'inp')
pr['A18'] = 'Pilot performance score (metrics at target)'; pr['F18'] = '=COUNTIF(G5:G16,"Yes")/COUNTA(A5:A16)'; pr['F18'].number_format = '0.0%'
pr['A19'] = 'Pilot goal'; pr['F19'] = 0.8; pr['F19'].number_format = '0%'; pr['F19'].font = inp; pr['F19'].fill = yellow
pr['A20'] = 'Goal met?'; pr['F20'] = '=IF(F18>=F19,"Yes","No")'
pr['A21'] = 'SLA checkpoints on time'; pr['F21'] = '=SUMIF(H5:H16,"Yes",E5:E16)'
pr['A22'] = 'SLA checkpoints total'; pr['F22'] = '=SUMIF(H5:H16,"Yes",D5:D16)'
pr['A23'] = 'Overall SLA compliance'; pr['F23'] = '=F21/F22'; pr['F23'].number_format = '0.0%'
for r in range(18, 24): pr.cell(r, 1).font = bold; pr.cell(r, 6).font = bold if r != 19 else inp
flag(pr, 'G5:G16', 'Yes', 'No'); widths(pr, [8, 52, 10, 12, 16, 10, 12, 22])

# ---------- KPI_Dashboard
kd = wb.create_sheet('KPI_Dashboard')
kd['A1'] = 'KPI Dashboard'; kd['A1'].font = title
kd['A2'] = 'All values are formulas linked to the other sheets, except blue inputs. Status updates automatically.'; kd['A2'].font = note
head(kd, 4, ['KPI', 'Baseline', 'Target', 'Current', 'Status', 'Source'])
k = [
 ('Time-to-hire (days)', '=Hires_Sample!E16', 25, 23, '=IF(D5<=C5,"Met","Not met")', 'Baseline: hires sample. Current: Week 2 dry run (input).', '0.0'),
 ('90-day retention', '=Hires_Sample!E17', 0.9, 'Pending', '=IF(ISNUMBER(D6),IF(D6>=C6,"Met","Not met"),"Pending")', 'Needs first cohort to reach Day 90.', '0%'),
 ('Feedback turnaround, median hours', 96, 24, '=MEDIAN(Scorecard!N5:N40)', '=IF(D7<=C7,"Met","Not met")', 'Baseline = 4 days (midpoint of est. 3 to 5, assumption).', '0.0'),
 ('Scorecards within SLA', 'Not tracked', 0.9, '=COUNTIF(Scorecard!O5:O40,"On time")/COUNTA(Scorecard!B5:B40)', '=IF(D8>=C8,"Met","Not met")', 'Scorecard sheet.', '0%'),
 ('Candidate satisfaction (CSAT)', 'Not measured', 0.85, '=Pilot_Results!F16', '=IF(D9>=C9,"Met","Not met")', 'Pilot survey, 20 responses.', '0%'),
 ('Onboarding LMS completion', 'Not measured', 1, '=Onboarding!J16', '=IF(D10>=C10,"Met",IF(D10>=0.9,"On track","Behind"))', 'Onboarding sheet; floor 90% before Day 14.', '0.0%'),
 ('Pilot performance score', 'n/a', 0.8, '=Pilot_Results!F18', '=IF(D11>=C11,"Met","Not met")', 'Pilot_Results sheet.', '0.0%'),
 ('Overall SLA compliance', 'n/a', 0.9, '=Pilot_Results!F23', '=IF(D12>=C12,"Met","Not met")', 'Target 90% is a team goal (assumption).', '0.0%'),
 ('Duplicate applications flagged', 'Not tracked', 'n/a', '=COUNTIF(Candidates!N5:N44,"Duplicate")', '=IF(D13>0,"Review","None")', 'Candidates sheet.', '0'),
]
for i, (a, b, c, d, e, f, fm) in enumerate(k):
    r = 5 + i
    put(kd, r, 1, a); x = put(kd, r, 2, b, 'calc' if str(b).startswith('=') else 'inp', fm if not isinstance(b, str) or b.startswith('=') else None)
    put(kd, r, 3, c, 'inp', fm if not isinstance(c, str) else None); put(kd, r, 4, d, 'calc' if str(d).startswith('=') else 'inp', fm if (not isinstance(d, str) or d.startswith('=')) else None)
    put(kd, r, 5, e); put(kd, r, 6, f); kd.cell(r, 6).font = note
for rr, fm in ((5, '0.0'), (6, '0%'), (7, '0')): kd.cell(rr, 2).number_format = fm
kd.conditional_formatting.add('E5:E13', CellIsRule(operator='equal', formula=['"Met"'], fill=green))
kd.conditional_formatting.add('E5:E13', CellIsRule(operator='equal', formula=['"Not met"'], fill=red))
kd.conditional_formatting.add('E5:E13', CellIsRule(operator='equal', formula=['"Behind"'], fill=red))
kd['A16'] = 'Pipeline funnel'; kd['A16'].font = bold
head(kd, 17, ['Stage', 'Candidates'])
for i, s in enumerate(stages):
    r = 18 + i; put(kd, r, 1, s); put(kd, r, 2, f'=COUNTIF(Candidates!$G$5:$G$44,A{r})')
put(kd, 25, 1, 'Total'); kd.cell(25, 1).font = bold; put(kd, 25, 2, '=SUM(B18:B24)'); kd.cell(25, 2).font = bold
widths(kd, [34, 14, 12, 14, 12, 56])
c1 = BarChart(); c1.type = 'bar'; c1.title = 'Candidates by stage'; c1.style = 10; c1.legend = None
c1.add_data(Reference(kd, min_col=2, min_row=17, max_row=24), titles_from_data=True); c1.set_categories(Reference(kd, min_col=1, min_row=18, max_row=24))
c1.height = 7.5; c1.width = 14; kd.add_chart(c1, 'D16')
c2 = BarChart(); c2.type = 'col'; c2.title = 'Pilot metrics: result vs target'; c2.style = 10
c2.add_data(Reference(pr, min_col=3, min_row=4, max_row=16), titles_from_data=True); c2.add_data(Reference(pr, min_col=6, min_row=4, max_row=16), titles_from_data=True)
c2.set_categories(Reference(pr, min_col=1, min_row=5, max_row=16)); c2.height = 8; c2.width = 22; c2.y_axis.number_format = '0%'; kd.add_chart(c2, 'A28')

wb.move_sheet('KPI_Dashboard', offset=-6)
for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            if c.font and c.font.name != FN: c.font = Font(name=FN, size=10, bold=c.font.b, italic=c.font.i, color=c.font.color)
wb.save('../data/HR_TA_Onboarding_Workbook.xlsx')
